# 大井競馬分析シート × 大井競馬分析表 を「日付＋R」で突合して、マザーデータ（hist.json）に新しいレースを足す
# 使い方：python3 -I merge_raw.py <分析シート.xlsx> <分析表.xlsx> [もとのhist.json] [出力hist.json]
#  - もとの280走（姫フィルターの検証データ）は書き換えない。食い違いがあれば報告だけする
#  - 新しいレース＝両方のファイルにそろっていて、もとのデータにない「日付＋R」
#  - 分析表とシートで着順がちがう → 分析表（オッズ側）を採用して備考に両方の値／同じ人気の重複 → 3着人気を空欄＋備考
import datetime, json, re, sys
import openpyxl

Z2H = str.maketrans('０１２３４５６７８９．', '0123456789.')
BABA = ('良', '稍重', '重', '不良')

def to_date(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, datetime.date): return v.isoformat()
    if isinstance(v, (int, float)) and 20000 < v < 80000: return (datetime.date(1899, 12, 30) + datetime.timedelta(days=round(v))).isoformat()
    if isinstance(v, str):
        s = v.translate(Z2H).strip()
        m = re.match(r'^(\d{4})[/\-.年](\d{1,2})[/\-.月](\d{1,2})', s)
        if m: return '%s-%02d-%02d' % (m.group(1), int(m.group(2)), int(m.group(3)))
    return None

def num(v):
    if v is None or isinstance(v, bool): return None
    if isinstance(v, (int, float)): return int(v) if float(v).is_integer() else float(v)
    s = str(v).translate(Z2H).replace(',', '').strip()
    try:
        f = float(s); return int(f) if f.is_integer() else f
    except ValueError:
        return None

def ranks(v):
    if v is None: return None
    xs = [int(x) for x in re.findall(r'\d+', str(v).translate(Z2H))]
    return xs[:3] if len(xs) >= 2 else None

def read_sheet(path):
    """分析シート：1行1レース（日付・R・馬場・距離・出馬数・結果）"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            c = list(row) + [None] * 6
            d = to_date(c[0]); r = num(c[1])
            if not d or not isinstance(r, int) or not 1 <= r <= 16: continue
            out[(d, r)] = {'baba': (str(c[2]).strip() if c[2] is not None else None), 'dist': num(c[3]), 'n': num(c[4]), 'res': c[5], 'rk': ranks(c[5])}
    return out

def read_table(path):
    """分析表：1行1日。2行目の「1R」「2R」…の列から、Rごとに 単勝1位・3連単1位・結果 を読む"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        cols = {}
        for r in rows[:6]:
            for i, v in enumerate(r):
                m = re.match(r'^\s*(\d{1,2})\s*R\s*$', str(v or ''))
                if m: cols[int(m.group(1))] = i
            if cols: break
        for r in rows:
            r = r + [None] * 60
            d = to_date(r[0])
            if not d: continue
            for R, i in cols.items():
                t1, s1, res = num(r[i]), num(r[i + 1]), r[i + 2]
                if t1 is None and s1 is None and res in (None, ''): continue
                out[(d, R)] = {'tan1': t1, 's1': s1, 'res': res, 'rk': ranks(res)}
    return out

def main(sheet_path, table_path, base_path='hist.json', out_path=None):
    base = json.load(open(base_path, encoding='utf8'))
    have = {(r['date'], r['r']): r for r in base}
    S, T = read_sheet(sheet_path), read_table(table_path)
    rep = {'sheet': len(S), 'table': len(T)}
    # もとのデータとのちがい（書き換えない・報告だけ）
    diffs = []
    for k, b in have.items():
        s, t = S.get(k), T.get(k)
        if not s or not t: diffs.append((k, '元ファイルに無い', bool(s), bool(t))); continue
        for f in ('baba', 'dist', 'n'):
            if s[f] != b.get(f): diffs.append((k, f, s[f], b.get(f)))
        for f in ('tan1', 's1'):
            if t[f] != b.get(f): diffs.append((k, f, t[f], b.get(f)))
        rk = t['rk'] or s['rk']
        want = [b.get('p1'), b.get('p2'), b.get('p3')]
        if rk and rk[:2] != want[:2]: diffs.append((k, '着順', rk, want))
    # 新しいレース
    new, skipped = [], []
    for k in sorted(set(S) | set(T)):
        if k in have: continue
        s, t = S.get(k), T.get(k)
        d, R = k
        if not s or not t:
            skipped.append((k, '分析シートだけ' if s else '分析表だけ')); continue
        if s['baba'] not in BABA or not s['dist'] or not s['n'] or not t['tan1'] or not t['s1'] or not (t['rk'] or s['rk']):
            skipped.append((k, '項目が足りない')); continue
        notes, rk = [], t['rk'] or s['rk']
        if t['rk'] and s['rk'] and t['rk'] != s['rk']:
            for j, nm in enumerate(['1着', '2着', '3着']):
                a = t['rk'][j] if j < len(t['rk']) else None; b2 = s['rk'][j] if j < len(s['rk']) else None
                if a != b2: notes.append(f'分析表とシートで{nm}人気不一致（表={a}・シート={b2}）。表を採用')
        p = rk + [None] * (3 - len(rk))
        if p[2] is not None and p[2] in (p[0], p[1]):
            notes.append(f'結果が{"-".join(str(x) + "位" for x in rk)}で3着人気が重複。3着人気は空欄（3着の集計から除外）'); p[2] = None
        rec = {'id': d.replace('-', '') + '-oi-%02d' % R, 'date': d, 'venue': '大井', 'r': R, 'baba': s['baba'], 'dist': s['dist'], 'surface': 'ダ', 'n': s['n'],
               'tan1': t['tan1'], 's1': t['s1'], 'p1': p[0], 'p2': p[1], 'p3': p[2], 'cnt': 1, 'note': '・'.join(notes)}
        # 念のための確認（おかしな値は報告）
        warn = []
        if not (1 <= rec['tan1'] < 100): warn.append('単勝1位')
        if rec['s1'] < rec['tan1']: warn.append('3連単1位<単勝1位')
        if not (5 <= rec['n'] <= 16): warn.append('頭数')
        if any(x and x > rec['n'] for x in p): warn.append('着順の人気が頭数より大きい')
        if len(set(x for x in p[:2])) < 2: warn.append('1着と2着が同じ')
        rec['_warn'] = warn
        new.append(rec)
    rep['new'] = len(new); rep['diffs'] = diffs; rep['skipped'] = skipped; rep['summary'] = summarize(new)
    if out_path:
        out = base + [{k: v for k, v in r.items() if k != '_warn'} for r in new]
        out.sort(key=lambda r: (r['date'], r['r'], r['id']))
        json.dump(out, open(out_path, 'w', encoding='utf8'), ensure_ascii=False)
        rep['total'] = len(out)
        if new:   # 更新の記録（Excelの「突合メモ」に出す）
            import os
            lp = os.path.join(os.path.dirname(os.path.abspath(out_path)), 'mother_log.json')
            log = json.load(open(lp, encoding='utf8')) if os.path.exists(lp) else []
            log.append({'date': datetime.date.today().isoformat(), 'text': f'{rep["summary"]}を足した（分析シート・分析表の最新版から）。合計{len(out)}走'})
            json.dump(log, open(lp, 'w', encoding='utf8'), ensure_ascii=False, indent=1)
    return rep, new

def summarize(new):
    """10/8（1〜12R）・10/9（1〜10R）の22走 のような短い説明"""
    by = {}
    for r in new: by.setdefault(r['date'], []).append(r['r'])
    parts = []
    for d, rs in sorted(by.items()):
        rs = sorted(rs); md = '%d/%d' % (int(d[5:7]), int(d[8:10]))
        if rs == list(range(rs[0], rs[-1] + 1)) and len(rs) > 1: parts.append(f'{md}（{rs[0]}〜{rs[-1]}R）')
        else: parts.append(f'{md}（{"・".join(str(x) for x in rs)}R）')
    return '・'.join(parts) + f'の{len(new)}走'

if __name__ == '__main__':
    a = sys.argv[1:]
    rep, new = main(*a)
    print('分析シート', rep['sheet'], '走／分析表', rep['table'], '走／新しいレース', rep['new'], '走', '→ 合計', rep.get('total', '-'), '|', rep['summary'])
    print('もとのデータとのちがい:', len(rep['diffs']))
    for d in rep['diffs'][:30]: print('   ', d)
    print('入れなかったレース（新しい日付のみ）:', [x for x in rep['skipped'] if x[0][0] > '2026-10-07'])
    print('入れなかったレース（〜10/7）:', len([x for x in rep['skipped'] if x[0][0] <= '2026-10-07']), '走', sorted(set(x[0][0] for x in rep['skipped'] if x[0][0] <= '2026-10-07')))
    for r in new:
        print('  +', r['date'], '%2dR' % r['r'], r['baba'], r['dist'], '%2d頭' % r['n'], '単勝', r['tan1'], '3連単', r['s1'], '着順', r['p1'], r['p2'], r['p3'], r['note'], ('⚠️' + ','.join(r['_warn'])) if r['_warn'] else '')
