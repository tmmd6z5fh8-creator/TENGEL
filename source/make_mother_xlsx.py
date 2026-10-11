# 大井競馬分析シート × 大井競馬分析表（日付＋Rで突合済み＝hist.json）→ アプリで読めるExcel（検証データ A〜P列）を1つ作る
# 使い方：python3 make_mother_xlsx.py [hist.json] [出力.xlsx]
# 形は姫フィルターの「検証データ」と同じ（1行目タイトル・2行目説明・3行目見出し・4行目からデータ）。アプリは「検証データ」シートのA〜P列を読む
import datetime, json, sys
from collections import Counter
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

SRC = sys.argv[1] if len(sys.argv) > 1 else 'hist.json'
OUT = sys.argv[2] if len(sys.argv) > 2 else '大井競馬_マザーデータ_アプリ用.xlsx'
rows = sorted(json.load(open(SRC, encoding='utf8')), key=lambda r: (r['date'], r['r'], r['id']))
FIRST, BLANK_TO = 4, 1000            # データは4行目から。新しいレース用に1000行目まで黄色の入力欄を用意
LAST = FIRST + len(rows) - 1

F = 'Arial'
font = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True)
title_font = Font(name=F, size=13, bold=True, color='16203D'); head_font = Font(name=F, size=10, bold=True, color='FFFFFF')
note_font = Font(name=F, size=9, color='5A6175')
yellow = PatternFill('solid', start_color='FFF6C8'); navy = PatternFill('solid', start_color='16203D'); grey = PatternFill('solid', start_color='EEF0F4')
thin = Side(style='thin', color='D7DBE4'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center'); wrap = Alignment(wrap_text=True, vertical='top')

HEAD = ['日付', 'R', '馬場', '距離(m)', '頭数', '単勝1位', '3連単1位', '1着人気', '2着人気', '3着人気', '3連単5位', '3連単10位', '馬単払戻(円)', '3連単払戻(円)', '備考', '集計対象']
KEYS = ['date', 'r', 'baba', 'dist', 'n', 'tan1', 's1', 'p1', 'p2', 'p3', 's5', 's10', 'umatanPay', 'sanrenPay', 'note', 'cnt']
FMT = {1: 'yyyy/m/d', 2: '0', 4: '0', 5: '0', 6: '0.0', 7: '0.0', 8: '0', 9: '0', 10: '0', 11: '0.0', 12: '0.0', 13: '#,##0', 14: '#,##0', 16: '0'}
WIDTH = {1: 11, 2: 5, 3: 6, 4: 8, 5: 6, 6: 8, 7: 9, 8: 8, 9: 8, 10: 8, 11: 9, 12: 9, 13: 11, 14: 11, 15: 46, 16: 8}

wb = openpyxl.Workbook()
# ---------- 検証データ ----------
ws = wb.active; ws.title = '検証データ'
ws['A1'] = '大井競馬　マザーデータ（検証データ）— テンゲルⅢアプリ・姫フィルター共通'; ws['A1'].font = title_font
ws['A2'] = ('黄色のセル（A〜P列）に1行1レースで入れてね。着順は「1着人気・2着人気・3着人気」の3つの列に人気の数字を別々に入れる（1位-2位-5位 → 1・2・5）。'
            '3連単5位〜3連単払戻・備考は空でもOK。集計対象はふだん1（大井以外・除外するレースは0）。')
ws['A2'].font = note_font; ws['A2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells('A2:P2'); ws.row_dimensions[2].height = 30
for c, h in enumerate(HEAD, 1):
    cell = ws.cell(3, c, h); cell.font = head_font; cell.fill = navy; cell.alignment = center; cell.border = box
for i, r in enumerate(rows):
    for c, k in enumerate(KEYS, 1):
        v = r.get(k)
        if k == 'date': v = datetime.datetime.strptime(v, '%Y-%m-%d')
        if k == 'note': v = v or None
        if k == 'cnt': v = 1 if v is None else v
        cell = ws.cell(FIRST + i, c, v)
for rr in range(FIRST, BLANK_TO + 1):
    for c in range(1, 17):
        cell = ws.cell(rr, c); cell.font = font; cell.fill = yellow; cell.border = box
        if c in FMT: cell.number_format = FMT[c]
        cell.alignment = Alignment(horizontal='left' if c == 15 else 'center', vertical='center')
for c, w in WIDTH.items(): ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
ws.freeze_panes = 'B4'
ws.cell(3, 3).comment = Comment('良・稍重・重・不良 から選んでね', '姫')
ws.cell(3, 8).comment = Comment('1着の馬の「単勝人気」の順位（例：1番人気なら1）。馬番ではないよ', '姫')
ws.cell(3, 16).comment = Comment('1＝集計する／0＝集計しない（大井以外・重複などで外すレース）', '姫')
ws.cell(3, 13).comment = Comment('当たり外れに関係なく、そのレースの馬単の払戻（100円あたり）。回収率の検証に使うよ', '姫')
rng = lambda col: f'{col}{FIRST}:{col}{BLANK_TO}'
for formula, col, msg in [('"良,稍重,重,不良"', 'C', '良・稍重・重・不良のどれかだよ'), ('"1,0"', 'P', '1か0だよ')]:
    dv = DataValidation(type='list', formula1=formula, allow_blank=True, error=msg, errorTitle='入力の確認'); dv.add(rng(col)); ws.add_data_validation(dv)
for col, lo, hi, msg in [('B', 1, 16, 'Rは1〜16だよ'), ('E', 1, 18, '頭数は1〜18だよ'), ('H', 1, 18, '人気は1〜18だよ'), ('I', 1, 18, '人気は1〜18だよ'), ('J', 1, 18, '人気は1〜18だよ')]:
    dv = DataValidation(type='whole', operator='between', formula1=str(lo), formula2=str(hi), allow_blank=True, error=msg, errorTitle='入力の確認'); dv.add(rng(col)); ws.add_data_validation(dv)

# ---------- 使い方 ----------
us = wb.create_sheet('使い方')
D = f"検証データ!$A${FIRST}:$A${BLANK_TO}"
lines = [
    ('テンゲルⅢ　マザーデータの使い方', title_font),
    ('', None),
    ('■ アプリに入れる（iPhone）', bold),
    ('1. このファイルをiPhoneの「ファイル」アプリに保存する（LINEで受け取ったときは、共有ボタン →「"ファイル"に保存」）', font),
    ('2. テンゲルⅢアプリ →「設定」タブ →「マザーデータ」の1つ目「大井競馬」の「編集」', font),
    ('3. 「Excelファイルを選ぶ」→ このファイルを選ぶ → 内容を確かめて「「大井競馬」に入れる」', font),
    ('   入れたあとでも「ひとつ前のデータに戻す」で前のデータに戻せるよ', note_font),
    ('', None),
    ('■ 新しいレースを足す', bold),
    ('「検証データ」シートの一番下の黄色い行に、1行1レースで入れる（分析シートと分析表の2つに分けなくてOK）', font),
    ('入れ方の例（この行は例なので、検証データには入っていないよ）↓', note_font),
]
for i, (t, f) in enumerate(lines, 1):
    c = us.cell(i, 1, t)
    if f: c.font = f
ex_row = len(lines) + 1
for c, h in enumerate(HEAD, 1):
    cell = us.cell(ex_row, c, h); cell.font = head_font; cell.fill = navy; cell.alignment = center; cell.border = box
example = [datetime.datetime(2026, 10, 19), 8, '稍重', 1400, 14, 2.4, 41.3, 1, 7, 2, 58.2, 84.0, 3250, None, '例：3連単払戻は空でもOK', 1]
for c, v in enumerate(example, 1):
    cell = us.cell(ex_row + 1, c, v); cell.font = font; cell.fill = grey; cell.border = box; cell.alignment = center
    if c in FMT: cell.number_format = FMT[c]
tips = [
    ('', None),
    ('■ ポイント', bold),
    ('・着順は「1着人気・2着人気・3着人気」に人気の順位を別々に入れる（分析シートの「1位-7位-2位」なら 1・7・2）', font),
    ('・同じ人気が重なる記載ミス（例 10位-4位-4位）は、3着人気を空欄にして備考に書く', font),
    ('・分析表と分析シートで着順がちがうときは、オッズ側（分析表）を使って、備考に両方の値を書く', font),
    ('・大井以外のレースは集計対象を0にして、備考に開催場と芝・ダートを書く', font),
    ('・馬単払戻は当たり外れに関係なく毎レース入れると、回収率の検証ができるよ（空でもアプリは読めるよ）', font),
    ('', None),
    ('■ いまのデータ（下の数字は検証データから自動で数えるよ）', bold),
]
base = ex_row + 2
for i, (t, f) in enumerate(tips):
    c = us.cell(base + i, 1, t)
    if f: c.font = f
k = base + len(tips)
stats = [
    ('収録レース', f'=COUNTIF({D},">0")', '走'),
    ('集計対象（P列＝1）', f'=COUNTIFS({D},">0",検証データ!$P${FIRST}:$P${BLANK_TO},1)', '走'),
    ('最初の日付', f'=MIN({D})', ''),
    ('最後の日付', f'=MAX({D})', ''),
]
for i, (lab, fml, unit) in enumerate(stats):
    us.cell(k + i, 1, lab).font = font
    c = us.cell(k + i, 2, fml); c.font = bold; c.number_format = 'yyyy/m/d' if '日付' in lab else '0'
    us.cell(k + i, 3, unit).font = font
us.column_dimensions['A'].width = 24
for c in range(2, 17): us.column_dimensions[openpyxl.utils.get_column_letter(c)].width = WIDTH.get(c, 10)
us.column_dimensions['B'].width = 12   # 下の「いまのデータ」の日付が入る幅

# ---------- 突合メモ ----------
ms = wb.create_sheet('突合メモ')
days = sorted(set(r['date'] for r in rows))
notes = [r for r in rows if r.get('note')]
memo = [
    ('このファイルの作り方（姫が作成・2026/10/11）', title_font),
    ('', None),
    ('■ 元データ', bold),
    ('大井競馬分析シート.xlsx：日付・R・馬場・距離・出馬数・結果（着順の人気）', font),
    ('大井競馬分析表.xlsx：日付ごとに、Rごとの単勝1位オッズ・3連単1位オッズ・結果', font),
    ('この2つを「日付＋R」で突き合わせて、1行1レースにまとめた（姫フィルター第13改訂版の検証データA〜P列と同じ中身）', font),
    ('', None),
    ('■ 収録', bold),
    (f'{len(rows)}走（{days[0][5:].replace("-", "/")}〜{days[-1][5:].replace("-", "/")}・{len(days)}日）。集計対象 {sum(1 for r in rows if r.get("cnt", 1) == 1)}走', font),
    ('', None),
    ('■ 入れていないレース', bold),
    ('7/1〜7/3の36走：分析シートにはあるが、分析表に単勝1位・3連単1位がないため（判定に使えない）', font),
    ('7/24 5R・6R：どちらのファイルにも距離・結果・オッズが入っていないため', font),
    ('', None),
    ('■ 集計から外したレース・食い違い（検証データの備考にも書いてあるよ）', bold),
]
for i, (t, f) in enumerate(memo, 1):
    c = ms.cell(i, 1, t)
    if f: c.font = f
hdr = len(memo) + 1
for c, h in enumerate(['日付', 'R', '集計対象', '内容'], 1):
    cell = ms.cell(hdr, c, h); cell.font = head_font; cell.fill = navy; cell.alignment = center; cell.border = box
grouped = Counter(r['note'] for r in notes if '重複転記' in r['note'])
line = hdr + 1
for r in notes:
    if '重複転記' in r['note']: continue
    vals = [datetime.datetime.strptime(r['date'], '%Y-%m-%d'), r['r'], r.get('cnt', 1), r['note']]
    for c, v in enumerate(vals, 1):
        cell = ms.cell(line, c, v); cell.font = font; cell.border = box
        if c == 1: cell.number_format = 'yyyy/m/d'
    line += 1
dup = [r for r in notes if '重複転記' in r['note']]
if dup:
    vals = [datetime.datetime.strptime(dup[0]['date'], '%Y-%m-%d'), f'1〜{max(r["r"] for r in dup)}R（{len(dup)}走）', 0, '9/2は、1〜2Rが9/3、3〜12Rが9/1と距離・頭数・結果まで同じで、重複転記の疑い。集計対象0（データは残す）']
    for c, v in enumerate(vals, 1):
        cell = ms.cell(line, c, v); cell.font = font; cell.border = box
        if c == 1: cell.number_format = 'yyyy/m/d'
    line += 1
ms.column_dimensions['A'].width = 12; ms.column_dimensions['B'].width = 14; ms.column_dimensions['C'].width = 9; ms.column_dimensions['D'].width = 90
for wsx in (us, ms):
    for row in wsx.iter_rows():
        for cell in row:
            if cell.value is not None and cell.font and cell.font.name != F:
                cell.font = Font(name=F, size=cell.font.size or 10, bold=cell.font.bold, color=cell.font.color)
wb.active = 0
wb.save(OUT)
print('saved', OUT, len(rows), 'rows', days[0], '..', days[-1])
