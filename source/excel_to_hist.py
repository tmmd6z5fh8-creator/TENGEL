# まさちゃんのExcel（姫フィルター）の「検証データ」A〜P列 → アプリに入れる基盤データ（hist.json）
# 使い方：python3 excel_to_hist.py 大井競馬_姫フィルター_第13改訂版.xlsx [出力先=hist.json]
# 読み方はiPhoneアプリの「設定 → 基盤データ → Excelファイルを選ぶ」と同じ（pwa_ui.js の rowToRec）
import datetime, json, re, sys
import openpyxl

BABA = ['良', '稍重', '重', '不良']
Z2H = str.maketrans('０１２３４５６７８９．／－', '0123456789./-')

def num(v):
    if v is None or isinstance(v, bool): return None
    if isinstance(v, (int, float)): return v
    s = str(v).translate(Z2H).replace(',', '').replace('円', '').strip()
    if s in ('', '-', '—'): return None
    try:
        f = float(s); return int(f) if f.is_integer() else f
    except ValueError:
        return None

def date_of(v):
    if v is None or v == '': return None
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, datetime.date): return v.isoformat()
    if isinstance(v, (int, float)):
        return (datetime.date(1899, 12, 30) + datetime.timedelta(days=round(v))).isoformat() if 20000 < v < 80000 else None
    s = str(v).translate(Z2H).strip()
    m = re.match(r'^(\d{4})[/\-.年](\d{1,2})[/\-.月](\d{1,2})', s)
    if m: return '%s-%02d-%02d' % (m.group(1), int(m.group(2)), int(m.group(3)))
    if re.match(r'^\d{5}(\.\d+)?$', s): return date_of(float(s))
    return None

def row_to_rec(c):
    date = date_of(c[0])
    if not date: return None
    r, baba = num(c[1]), str(c[2] or '').strip()
    if not r or r < 1 or r > 16 or r != int(r) or baba not in BABA: return 'bad'
    r = int(r); cnt = num(c[15])
    def n(i):
        x = num(c[i]); return int(x) if isinstance(x, float) and x.is_integer() else x
    return {'id': date.replace('-', '') + '-oi-%02d' % r, 'date': date, 'venue': '大井', 'r': r, 'baba': baba, 'dist': n(3), 'surface': 'ダ', 'n': n(4),
            'tan1': n(5), 's1': n(6), 'p1': n(7), 'p2': n(8), 'p3': n(9), 's5': n(10), 's10': n(11), 'umatanPay': n(12), 'sanrenPay': n(13),
            'note': '' if c[14] is None else str(c[14]).strip(), 'cnt': 1 if cnt is None else cnt}

def main(path, out='hist.json'):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    names = [n for n in wb.sheetnames if '検証データ' in n] or wb.sheetnames
    best = None
    for name in names:
        recs, bad, seen = [], 0, {}
        for row in wb[name].iter_rows(min_col=1, max_col=16, values_only=True):
            c = list(row) + [None] * (16 - len(row))
            x = row_to_rec(c)
            if x == 'bad': bad += 1
            elif x: recs.append(x)
        for r in recs:
            if r['id'] in seen: seen[r['id']] += 1; r['id'] += '-%d' % seen[r['id']]
            else: seen[r['id']] = 1
        if recs and (not best or len(recs) > len(best[1])): best = (name, recs, bad)
        if best and '検証データ' in name: break
    if not best: sys.exit('レースの行が見つからなかったよ')
    name, recs, bad = best
    # アプリに入れるのは、これまでの hist.json と同じ項目（払戻などは入っているときだけ）
    slim = []
    for r in recs:
        o = {k: r[k] for k in ['id', 'date', 'venue', 'r', 'baba', 'dist', 'surface', 'n', 'tan1', 's1', 'p1', 'p2', 'p3', 'cnt', 'note']}
        for k in ['s5', 's10', 'umatanPay', 'sanrenPay']:
            if r[k] is not None: o[k] = r[k]
        slim.append(o)
    json.dump(slim, open(out, 'w', encoding='utf8'), ensure_ascii=False)
    ds = sorted(r['date'] for r in slim)
    use = sum(1 for r in slim if r['cnt'] == 1 and r['p1'] and r['p2'])
    print(f'{name}: {len(slim)}走（{ds[0]}〜{ds[-1]}）・集計対象 {use}走・読めない行 {bad} → {out}')
    print('MOTHER_VER の目安:', ds[-1])

if __name__ == '__main__':
    main(*sys.argv[1:])
